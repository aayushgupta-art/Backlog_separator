import zipfile
from pathlib import Path
import openpyxl
import pymupdf
import pytest
from config import get_default_settings
from data_parser import parse_pdf
from course_grouper import build_exams
from scheduler import schedule_exams
from verifier import verify
from exporters import to_excel, to_csv, build_zip, export_all

REAL_PDF_PATH = Path("sample_input/Backlog_List.pdf")
SKIP_IF_NO_REAL_PDF = pytest.mark.skipif(
    not REAL_PDF_PATH.is_file(),
    reason="sample_input/Backlog_List.pdf not present"
)

@SKIP_IF_NO_REAL_PDF
def test_exporters_full_flow():
    assert REAL_PDF_PATH.is_file()
    parse_res = parse_pdf(REAL_PDF_PATH)
    exams, duplicates = build_exams(parse_res.rows)
    settings = get_default_settings()
    schedule = schedule_exams(exams, settings)
    v_res = verify(parse_res, None, schedule, settings)

    out_dir = Path("output")
    exports = export_all(parse_res, None, schedule, settings, out_dir)

    excel_p = exports["excel_path"]
    assert excel_p.is_file()

    # Verify Excel workbook with openpyxl
    wb = openpyxl.load_workbook(excel_p)
    assert set(wb.sheetnames) == {"Timetable", "Exam-wise students", "Student-wise schedule"}

    sheet_tt = wb["Timetable"]
    assert sheet_tt.max_row == len(exams) + 1

    sheet_exam = wb["Exam-wise students"]
    assert sheet_exam.max_row > len(exams)

    sheet_stu = wb["Student-wise schedule"]
    assert sheet_stu.max_row > len(exams)

    # Verify CSV file
    csv_p = exports["csv_path"]
    assert csv_p.is_file()
    assert csv_p.stat().st_size > 0

    # Verify ZIP file
    zip_p = exports["zip_path"]
    assert zip_p.is_file()

    with zipfile.ZipFile(zip_p, "r") as zf:
        members = zf.namelist()
        print("\nZIP Members:", members)
        assert "Audit_Report.pdf" in members
        assert "Timetable.pdf" in members
        assert "Attendance_Sheets.pdf" in members
        assert "Hall_Tickets.pdf" in members
        assert "Complete_Report.pdf" in members
        assert "Backlog_Schedule.xlsx" in members
        assert "Backlog_Schedule.csv" in members

    used_dates = set(it.slot.date for it in schedule.items)

    print("\n--- PHASE 7 EXPORT SUMMARY ---")
    print(f"Total Rows: {v_res.rows_read}")
    print(f"Unique Students: {v_res.unique_students}")
    print(f"Exams: {v_res.exams}")
    print(f"Duplicate Rows: {v_res.duplicate_rows}")
    print(f"Days Used: {len(used_dates)}")
    print(f"Slots Used: {schedule.slots_used}")
    print(f"Verify OK: {v_res.ok}")

    files_to_report = [
        exports["audit_pdf"],
        exports["timetable_pdf"],
        exports["attendance_pdf"],
        exports["hall_tickets_pdf"],
        exports["complete_report_pdf"],
        exports["excel_path"],
        exports["csv_path"],
        exports["zip_path"],
    ]

    print("\n--- FILE DETAILS ---")
    for f in files_to_report:
        size_kb = f.stat().st_size / 1024
        page_info = ""
        if f.suffix.lower() == ".pdf":
            doc = pymupdf.open(str(f))
            page_info = f"({len(doc)} pages)"
            doc.close()
        print(f"{f.name}: {size_kb:.1f} KB {page_info}")
